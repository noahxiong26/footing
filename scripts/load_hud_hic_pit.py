"""Fill footing/static/coc_capacity.csv from HUD's AHAR Part 1 downloads.

Download both files from https://www.huduser.gov/portal/node/11318 (2024 AHAR Part 1):
  * "HIC by CoC" (2007-2024) .xlsx
  * "PIT estimates by CoC" (2007-2024) .xlsb

Then run:
  python scripts/load_hud_hic_pit.py --hic 2007-2024-HIC-Counts-by-CoC.xlsx \
      --pit 2007-2024-PIT-Counts-by-CoC.xlsb --year 2024

Each workbook has one sheet per year. Column names have shifted between years, so columns are
matched by keywords rather than exact names, and the script prints which column it used for
each value so a person can check it. NOT yet run against the real files (the build sandbox could
not download them); verify the printed column choices on first use.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

STATIC = Path(__file__).resolve().parents[1] / "footing" / "static"

WANT = {
    "coc": [r"^coc number$", r"^coc$", r"coc number"],
    "es_beds": [r"total year-round beds \(es\)", r"year.?round beds.*\bes\b"],
    "th_beds": [r"total year-round beds \(th\)", r"year.?round beds.*\bth\b"],
    "sh_beds": [r"total year-round beds \(sh\)", r"year.?round beds.*\bsh\b"],
    "pit_total": [r"^overall homeless$", r"overall homeless"],
    "pit_unsheltered": [r"^unsheltered homeless$", r"unsheltered homeless"],
}


def pick(headers: list[str], key: str) -> int | None:
    norm = [str(h or "").strip().lower() for h in headers]
    for pat in WANT[key]:
        for i, h in enumerate(norm):
            if re.search(pat, h):
                return i
    return None


def read_sheet(path: Path, year: int) -> tuple[list[str], list[list]]:
    if path.suffix == ".xlsb":
        from pyxlsb import open_workbook
        with open_workbook(str(path)) as wb:
            name = next((s for s in wb.sheets if str(year) in s), None)
            if not name:
                sys.exit(f"{path.name}: no sheet named for {year}; sheets are {wb.sheets}")
            with wb.get_sheet(name) as sh:
                rows = [[c.v for c in r] for r in sh.rows()]
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        name = next((s for s in wb.sheetnames if str(year) in s), None)
        if not name:
            sys.exit(f"{path.name}: no sheet named for {year}; sheets are {wb.sheetnames}")
        rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    # header = first row that contains a CoC column
    for i, r in enumerate(rows[:10]):
        if pick([str(c) for c in r], "coc") is not None:
            return [str(c) for c in r], rows[i + 1:]
    sys.exit(f"{path.name}/{name}: could not find a header row with a CoC Number column")


def index(path: Path, year: int, keys: list[str]) -> dict[str, dict]:
    headers, rows = read_sheet(path, year)
    cols = {k: pick(headers, k) for k in ["coc", *keys]}
    for k, i in cols.items():
        print(f"  {path.name}: {k:<16} -> {headers[i] if i is not None else 'NOT FOUND'}")
    out = {}
    for r in rows:
        coc = r[cols["coc"]] if cols["coc"] is not None else None
        if not coc:
            continue
        out[str(coc).strip()] = {k: (r[i] if i is not None else None) for k, i in cols.items() if k != "coc"}
    return out


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--hic", type=Path, required=True)
    ap.add_argument("--pit", type=Path, required=True)
    ap.add_argument("--year", type=int, default=2024)
    ap.add_argument("--crosswalk", type=Path, default=STATIC / "county_coc.csv")
    ap.add_argument("--out", type=Path, default=STATIC / "coc_capacity.csv")
    args = ap.parse_args(argv)
    hic = index(args.hic, args.year, ["es_beds", "th_beds", "sh_beds"])
    pit = index(args.pit, args.year, ["pit_total", "pit_unsheltered"])
    with open(args.crosswalk, newline="", encoding="utf-8") as f:
        crosswalk = list(csv.DictReader(f))
    out_rows = []
    for row in crosswalk:
        coc = row["coc_code"]
        h, p = hic.get(coc), pit.get(coc)
        if not h or not p:
            print(f"  ! {coc} missing from {'HIC' if not h else 'PIT'} {args.year}; skipped")
            continue
        out_rows.append({"county_fips": row["county_fips"], "coc_code": coc, "year": args.year,
                         "es_beds": h["es_beds"] or 0, "th_beds": h["th_beds"] or 0, "sh_beds": h["sh_beds"] or 0,
                         "pit_unsheltered": p["pit_unsheltered"] or 0, "pit_total": p["pit_total"] or 0,
                         "source": f"HUD AHAR {args.year} HIC+PIT by CoC"})
    with open(args.out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["county_fips", "coc_code", "year", "es_beds", "th_beds", "sh_beds",
                                          "pit_unsheltered", "pit_total", "source"])
        w.writeheader()
        w.writerows(out_rows)
    print(f"wrote {len(out_rows)} county rows to {args.out}")


if __name__ == "__main__":
    main()
